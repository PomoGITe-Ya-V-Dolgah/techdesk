"""Импорт реестра оборудования из Excel (.xlsx) или CSV.

Первая строка — заголовки. Порядок колонок любой, лишние колонки игнорируются.
Если инвентарный номер уже есть в базе — карточка обновляется.
Шаблон: /equipment/import/template.csv
"""

import csv
import io
import re
from dataclasses import dataclass, field

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.validators import validate_ipv46_address
from django.db import transaction
from django.db.models import Q

from .models import Department, Equipment, EquipmentCategory, Location, Software, SoftwareInstallation

COLUMNS = {
    "inventory_number": ["инвентарный номер", "инв. номер", "инв номер", "инв.номер", "inventory_number"],
    "name": ["название", "наименование", "name"],
    "category": ["категория", "тип", "category"],
    "location": ["кабинет", "место", "место установки", "location"],
    "assigned_user": ["сотрудник", "пользователь", "за кем закреплено", "user"],
    "department": ["подразделение", "отдел", "department"],
    "status": ["статус", "status"],
    "criticality": ["критичность", "criticality"],
    "manufacturer": ["производитель", "manufacturer"],
    "model": ["модель", "model"],
    "serial_number": ["серийный номер", "s/n", "serial_number"],
    "hostname": ["имя пк", "имя компьютера", "имя компьютера в сети", "hostname"],
    "ip_address": ["ip", "ip-адрес", "ip адрес", "ip_address"],
    "mac_address": ["mac", "mac-адрес", "mac адрес", "mac_address"],
    "operating_system": ["ос", "операционная система", "operating_system"],
    "software": ["по", "программы", "установленное по", "software"],
    "parent": ["подключено к", "установлено на", "parent"],
    "notes": ["примечание", "комментарий", "notes"],
}
REQUIRED = ["inventory_number", "name", "category", "location"]
TEMPLATE_HEADER = [
    "Инв. номер", "Название", "Категория", "Кабинет", "Сотрудник", "Подразделение", "Статус",
    "Производитель", "Модель", "Серийный номер", "Имя ПК", "IP", "MAC", "ОС", "ПО", "Подключено к", "Примечание",
]
TEMPLATE_ROWS = [
    ["PC-205-014", "ПК бухгалтера", "Рабочая станция", "Кабинет 205", "Петрова Елена", "Бухгалтерия", "В работе",
     "Lenovo", "ThinkCentre M70", "SN12345", "BUH-01", "192.168.1.21", "", "Windows 11 Pro",
     "1С:Предприятие 8.3, Microsoft Office 2021, Kaspersky", "", ""],
    ["MON-205-014", "Монитор бухгалтера", "Монитор", "Кабинет 205", "Петрова Елена", "Бухгалтерия", "В работе",
     "Dell", "P2422H", "", "", "", "", "", "", "PC-205-014", ""],
    ["MFU-2-001", "МФУ общее, 2 этаж", "МФУ", "Коридор 2 этаж", "", "", "В работе",
     "Kyocera", "M2540dn", "", "MFU-2", "192.168.1.50", "", "", "", "", "Картридж TK-1170"],
]


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def _normalize(header: str) -> str:
    return re.sub(r"\s+", " ", (header or "").strip().lower().replace("ё", "е"))


def _map_headers(headers):
    aliases = {_normalize(alias): key for key, names in COLUMNS.items() for alias in names}
    return {index: aliases[_normalize(h)] for index, h in enumerate(headers) if _normalize(h) in aliases}


def read_rows(uploaded_name: str, data: bytes) -> list[list[str]]:
    if uploaded_name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if value is None else str(value).strip() for value in row])
        return rows
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1251")  # «CSV (разделители — точка с запятой)» из русского Excel
    first_line = text.splitlines()[0] if text else ""
    delimiter = max([";", ",", "\t"], key=first_line.count)
    return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _choice(value: str, choices, default):
    value = (value or "").strip().lower()
    if not value:
        return default
    for key, label in choices:
        if value in (key.lower(), label.lower()):
            return key
    raise ValidationError(f"неизвестное значение «{value}»")


def find_user(value: str):
    value = (value or "").strip()
    if not value:
        return None
    user = User.objects.filter(username__iexact=value).first()
    if user:
        return user
    words = value.split()
    query = User.objects.all()
    for word in words:
        query = query.filter(Q(first_name__iexact=word) | Q(last_name__iexact=word))
    return query.first() if len(words) >= 2 else None


def import_equipment(uploaded_name: str, data: bytes) -> ImportResult:
    result = ImportResult()
    rows = [row for row in read_rows(uploaded_name, data) if any(row)]
    if not rows:
        result.errors.append((0, "Файл пустой."))
        return result

    mapping = _map_headers(rows[0])
    missing = [key for key in REQUIRED if key not in mapping.values()]
    if missing:
        names = ", ".join(COLUMNS[key][0] for key in missing)
        result.errors.append((1, f"Нет обязательных колонок: {names}. Скачайте шаблон и сверьте заголовки."))
        return result

    parents = []
    for line_no, row in enumerate(rows[1:], start=2):
        values = {key: row[index] if index < len(row) else "" for index, key in mapping.items()}
        try:
            with transaction.atomic():
                created = _import_row(values, line_no, result)
        except ValidationError as error:
            result.errors.append((line_no, "; ".join(error.messages)))
            continue
        if created:
            result.created += 1
        else:
            result.updated += 1
        if values.get("parent"):
            parents.append((line_no, values["inventory_number"], values["parent"]))

    # Связи «подключено к» — после всех строк, чтобы порядок строк был не важен.
    for line_no, inventory, parent_inventory in parents:
        parent = Equipment.objects.filter(inventory_number=parent_inventory).first()
        if parent:
            Equipment.objects.filter(inventory_number=inventory).update(parent=parent)
        else:
            result.warnings.append((line_no, f"не найдено оборудование «{parent_inventory}» для «Подключено к»"))
    return result


def _import_row(values: dict, line_no: int, result: ImportResult) -> bool:
    for key in REQUIRED:
        if not values.get(key):
            raise ValidationError(f"не заполнено поле «{COLUMNS[key][0]}»")

    category, _ = EquipmentCategory.objects.get_or_create(name=values["category"])
    location = Location.objects.filter(name=values["location"]).first() or Location.objects.create(name=values["location"])
    department = None
    if values.get("department"):
        department, _ = Department.objects.get_or_create(name=values["department"])

    ip = values.get("ip_address") or None
    if ip:
        try:
            validate_ipv46_address(ip)
        except ValidationError:
            raise ValidationError(f"неверный IP-адрес «{ip}»")

    notes = values.get("notes", "")
    user = find_user(values.get("assigned_user"))
    if values.get("assigned_user") and not user:
        result.warnings.append(
            (line_no, f"сотрудник «{values['assigned_user']}» не найден среди пользователей — записан в примечание")
        )
        notes = f"Сотрудник: {values['assigned_user']}\n{notes}".strip()

    defaults = {
        "name": values["name"],
        "category": category,
        "location": location,
        "department": department,
        "assigned_user": user,
        "status": _choice(values.get("status"), Equipment.Status.choices, Equipment.Status.IN_USE),
        "criticality": _choice(values.get("criticality"), Equipment.Criticality.choices, Equipment.Criticality.MEDIUM),
        "ip_address": ip,
        "notes": notes,
    }
    for key in ("manufacturer", "model", "serial_number", "hostname", "mac_address", "operating_system"):
        defaults[key] = values.get(key, "")

    equipment, created = Equipment.objects.update_or_create(
        inventory_number=values["inventory_number"], defaults=defaults
    )

    for name in re.split(r"[,;\n]", values.get("software", "")):
        name = name.strip()
        if name:
            software, _ = Software.objects.get_or_create(name=name)
            SoftwareInstallation.objects.get_or_create(equipment=equipment, software=software)
    return created


def template_csv() -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(TEMPLATE_HEADER)
    writer.writerows(TEMPLATE_ROWS)
    return buffer.getvalue()
